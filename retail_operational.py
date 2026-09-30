"""Extração auditável de indicadores operacionais de Varejo.

Planilhas oficiais têm precedência. Markdown é aceito somente como derivado
de PDF oficial e preenche lacunas de métrica/período não cobertas por Excel.
Ausência, baixa confiança e conceitos incompatíveis permanecem nulos.
"""

from __future__ import annotations

import math
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from metric_definitions import RETAIL_METRIC_IDS, RETAIL_OPERATIONAL_DICTIONARY


RETAIL_SCHEMA_VERSION = "retail_operational_v1"
RETAIL_EXTRACTOR_VERSION = "retail_operational_v1"
MISSING_MARKERS = {"", "-", "--", "—", "n/a", "na", "nd", "n.d.", "não divulgado", "nao divulgado"}


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().lower()


def normalize_period(value: Any) -> str | None:
    text = normalize_text(value).upper().replace(" ", "").replace("Q", "T")
    match = re.search(r"([1-4])T(?:R?I?M?)?[-_/]?(20\d{2}|\d{2})", text)
    if match:
        return f"{match.group(1)}T{match.group(2)[-2:]}"
    match = re.search(r"(?:FY|ANO|ANUAL)?[-_/]?(20\d{2})", text)
    if match:
        return match.group(1)
    return None


def parse_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip()
    if normalize_text(text) in MISSING_MARKERS or "%" in text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^\d,.-]", "", text)
    if not cleaned or cleaned in {"-", ".", ","}:
        return None
    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif cleaned.count(".") > 1:
        cleaned = cleaned.replace(".", "")
    try:
        number = float(cleaned)
    except ValueError:
        return None
    number = -abs(number) if negative else number
    return number if math.isfinite(number) else None


def identify_metric(label: str, context: str = "") -> tuple[str | None, list[str]]:
    label_text = normalize_text(label)
    haystack = normalize_text(f"{label} {context}")
    flags: list[str] = []
    for metric_id in RETAIL_METRIC_IDS:
        definition = RETAIL_OPERATIONAL_DICTIONARY[metric_id]
        if definition.get("derived"):
            continue
        aliases = tuple(normalize_text(alias) for alias in definition.get("aliases", ()))
        if not any(alias and alias in label_text for alias in aliases):
            continue
        forbidden = next((term for term in definition.get("forbidden_contexts", ()) if normalize_text(term) in haystack), None)
        if forbidden:
            return None, [f"forbidden_context:{normalize_text(forbidden)}"]
        if metric_id == "stores_count" and any(term in label_text for term in ("propria", "franqueada", "franquia", "bandeira", "marca")) and "total" not in label_text:
            flags.append("breakdown_without_explicit_total")
        return metric_id, flags
    # Aberturas específicas por tipo são preservadas como detalhe, nunca
    # somadas automaticamente para produzir um total.
    if "loja" in label_text and any(term in label_text for term in ("propria", "franqueada", "franquia")):
        return "stores_count", ["breakdown_without_explicit_total"]
    return None, []


def normalize_value(metric_id: str, value: float, unit_context: str) -> tuple[float, str, str]:
    text = normalize_text(unit_context).replace("²", "2")
    original_unit = unit_context.strip() or RETAIL_OPERATIONAL_DICTIONARY[metric_id]["unit"]
    if metric_id == "sales_area_sqm":
        factor = 1_000.0 if re.search(r"\bmil(?:hares)?\s*(?:m2|metros)", text) else 1.0
        return value * factor, "m²", original_unit
    if metric_id in {"physical_revenue_brl", "digital_revenue_brl", "comparable_total_revenue_brl"}:
        if any(term in text for term in ("bilhao", "bilhoes", "billion", "r$ bi")):
            factor = 1_000_000_000.0
        elif any(term in text for term in ("milhao", "milhoes", "million", "r$ mm")):
            factor = 1_000_000.0
        elif re.search(r"\b(r\$|brl)?\s*mil(?:hares)?\b", text):
            factor = 1_000.0
        else:
            factor = 1.0
        return value * factor, "R$", original_unit
    return value, str(RETAIL_OPERATIONAL_DICTIONARY[metric_id]["unit"]), original_unit


def _segment_from_label(label: str) -> str:
    text = normalize_text(label)
    if "franque" in text or "franquia" in text:
        return "franqueadas"
    if "propria" in text:
        return "próprias"
    return "consolidado"


def build_observation(*, ticker: str, metric_id: str, value: float, period: str,
                      label: str, unit_context: str, source_document: str,
                      source_type: str, source_url: str = "", sheet: str = "",
                      source_cell: str = "", page: int | None = None,
                      flags: Iterable[str] = ()) -> dict[str, Any]:
    normalized_value, unit, raw_unit = normalize_value(metric_id, value, unit_context)
    validation_flags = list(flags)
    breakdown = "breakdown_without_explicit_total" in validation_flags
    area_basis = "average" if metric_id == "sales_area_sqm" and "media" in normalize_text(label) else "closing"
    confidence = "medium" if breakdown else "high"
    return {
        "sector": "varejo", "ticker": ticker, "indicator_id": metric_id,
        "indicator_name": RETAIL_OPERATIONAL_DICTIONARY[metric_id]["display_name"],
        "value": normalized_value, "unit": unit, "period": period,
        "period_type": "quarter" if re.fullmatch(r"[1-4]T\d{2}", period) else "annual",
        "nature": RETAIL_OPERATIONAL_DICTIONARY[metric_id]["nature"],
        "scope": "retail_channels", "segment": _segment_from_label(label),
        "consolidated_or_breakdown": "breakdown" if breakdown else "consolidated",
        "area_basis": area_basis if metric_id == "sales_area_sqm" else None,
        "reported_or_derived": "reported", "confidence": confidence,
        "validation_status": "valid", "validation_flags": validation_flags,
        "source_document": source_document, "source_url": source_url,
        "source_type": source_type, "collected_at": datetime.now(timezone.utc).isoformat(),
        "sheet": sheet or None, "page": page, "source_cell": source_cell or None,
        "row_label": label, "column_label": period, "raw_value": str(value),
        "raw_unit": raw_unit, "normalization_rule": f"{raw_unit or 'unit_unspecified'} -> {unit}",
    }


def _nearest_period(rows: list[list[Any]], row_index: int, column: int, default_period: str | None) -> str | None:
    for index in range(row_index, max(-1, row_index - 12), -1):
        if column < len(rows[index]):
            period = normalize_period(rows[index][column])
            if period:
                return period
    return default_period


def extract_rows_observations(rows: list[list[Any]], *, ticker: str, source_document: str,
                              source_type: str, source_url: str = "", sheet: str = "",
                              default_period: str | None = None, page: int | None = None) -> list[dict[str, Any]]:
    observations: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows):
        label_column = next((index for index, cell in enumerate(row) if isinstance(cell, str) and cell.strip()), None)
        if label_column is None:
            continue
        label = str(row[label_column]).strip()
        context_rows = rows[max(0, row_index - 3):row_index + 1]
        context = " ".join(str(cell) for context_row in context_rows for cell in context_row if cell not in (None, ""))
        metric_id, flags = identify_metric(label, context)
        if not metric_id:
            continue
        for column in range(label_column + 1, len(row)):
            number = parse_number(row[column])
            if number is None:
                continue
            period = _nearest_period(rows, row_index, column, default_period)
            if not period:
                continue
            unit_context = f"{label} {context}"
            observations.append(build_observation(
                ticker=ticker, metric_id=metric_id, value=number, period=period,
                label=label, unit_context=unit_context, source_document=source_document,
                source_type=source_type, source_url=source_url, sheet=sheet,
                source_cell=f"R{row_index + 1}C{column + 1}", page=page, flags=flags,
            ))
    return dedupe_observations(observations)


def _workbook_rows(path: Path) -> list[tuple[str, list[list[Any]]]]:
    if path.suffix.lower() == ".xls":
        import pandas as pd
        book = pd.ExcelFile(path)
        try:
            return [(name, pd.read_excel(book, sheet_name=name, header=None).where(lambda frame: frame.notna(), None).values.tolist()) for name in book.sheet_names]
        finally:
            book.close()
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]) for sheet in workbook.worksheets]
    finally:
        workbook.close()


def extract_workbook_observations(path: str | Path, *, ticker: str, source_url: str = "") -> list[dict[str, Any]]:
    workbook_path = Path(path)
    default_period = normalize_period(workbook_path.name)
    observations: list[dict[str, Any]] = []
    for sheet_name, rows in _workbook_rows(workbook_path):
        observations.extend(extract_rows_observations(
            rows, ticker=ticker, source_document=workbook_path.name,
            source_type="official_spreadsheet", source_url=source_url,
            sheet=sheet_name, default_period=default_period,
        ))
    return dedupe_observations(observations)


def extract_markdown_observations(text: str, *, ticker: str, source_document: str,
                                  source_url: str = "") -> list[dict[str, Any]]:
    lines = text.splitlines()
    default_period = normalize_period(source_document)
    observations: list[dict[str, Any]] = []
    page: int | None = None
    index = 0
    while index < len(lines):
        page_match = re.search(r"(?:pagina|página|page)\s+(\d+)", lines[index], re.I)
        if page_match:
            page = int(page_match.group(1))
        if not lines[index].lstrip().startswith("|"):
            index += 1
            continue
        table: list[list[Any]] = []
        while index < len(lines) and lines[index].lstrip().startswith("|"):
            cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                table.append(cells)
            index += 1
        observations.extend(extract_rows_observations(
            table, ticker=ticker, source_document=source_document,
            source_type="official_pdf_markdown", source_url=source_url,
            default_period=default_period, page=page,
        ))
    return dedupe_observations(observations)


def dedupe_observations(observations: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    selected: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    rank = {"official_spreadsheet": 3, "official_pdf_markdown": 2}
    for observation in observations:
        key = (
            str(observation.get("indicator_id")), str(observation.get("period")),
            str(observation.get("scope")), str(observation.get("segment")),
        )
        current = selected.get(key)
        candidate_rank = (rank.get(str(observation.get("source_type")), 0), observation.get("confidence") == "high")
        current_rank = (rank.get(str(current.get("source_type")), 0), current.get("confidence") == "high") if current else (-1, False)
        if current is None or candidate_rank > current_rank:
            selected[key] = observation
    return sorted(selected.values(), key=lambda item: (str(item.get("period")), str(item.get("indicator_id")), str(item.get("segment"))))


def derive_metrics(observations: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    source = [item for item in observations if item.get("validation_status") == "valid" and item.get("consolidated_or_breakdown") == "consolidated"]
    by_period: dict[str, dict[str, dict[str, Any]]] = {}
    for item in source:
        by_period.setdefault(str(item["period"]), {})[str(item["indicator_id"])] = item
    derived: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []
    for period, metrics in by_period.items():
        digital = metrics.get("digital_revenue_brl")
        total = metrics.get("comparable_total_revenue_brl")
        if digital and total and total["value"] not in (None, 0) and digital.get("scope") == total.get("scope"):
            derived.append(_derived_observation("digital_revenue_share", period, digital["value"] / total["value"], "%", digital, total))
        else:
            diagnostics.append({"indicator_id": "digital_revenue_share", "period": period, "status": "not_calculated", "reason": "missing_or_incompatible_comparable_revenue_base"})
        physical = metrics.get("physical_revenue_brl")
        area = metrics.get("sales_area_sqm")
        if physical and area and area["value"] not in (None, 0) and physical.get("scope") == area.get("scope"):
            item = _derived_observation("physical_revenue_per_sqm", period, physical["value"] / area["value"], "R$/m²", physical, area)
            item["area_basis"] = area.get("area_basis") or "closing"
            if item["area_basis"] == "closing":
                item["validation_flags"] = ["closing_area_used_instead_of_average_area"]
                item["confidence"] = "medium"
            derived.append(item)
        else:
            diagnostics.append({"indicator_id": "physical_revenue_per_sqm", "period": period, "status": "not_calculated", "reason": "missing_or_incompatible_physical_revenue_or_area"})
    return derived, diagnostics


def _derived_observation(metric_id: str, period: str, value: float, unit: str,
                         numerator: dict[str, Any], denominator: dict[str, Any]) -> dict[str, Any]:
    return {
        "sector": "varejo", "ticker": numerator["ticker"], "indicator_id": metric_id,
        "indicator_name": RETAIL_OPERATIONAL_DICTIONARY[metric_id]["display_name"],
        "value": value, "unit": unit, "period": period,
        "period_type": numerator.get("period_type"), "nature": "ratio",
        "scope": numerator.get("scope"), "segment": "consolidado",
        "consolidated_or_breakdown": "consolidated", "reported_or_derived": "derived",
        "confidence": "high", "validation_status": "valid", "validation_flags": [],
        "source_document": numerator.get("source_document"), "source_url": numerator.get("source_url"),
        "source_type": "calculation", "collected_at": datetime.now(timezone.utc).isoformat(),
        "formula": f"{metric_id}: numerator / denominator",
        "inputs": {"numerator": numerator["indicator_id"], "denominator": denominator["indicator_id"]},
        "normalization_rule": "same_period_same_scope_only",
    }


def build_metricas(observations: Iterable[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    metricas: dict[str, list[dict[str, Any]]] = {}
    grouped: dict[str, list[dict[str, Any]]] = {}
    for observation in observations:
        if observation.get("validation_status") != "valid" or observation.get("consolidated_or_breakdown") != "consolidated":
            continue
        grouped.setdefault(str(observation["indicator_id"]), []).append(observation)
    for metric_id, items in grouped.items():
        definition = RETAIL_OPERATIONAL_DICTIONARY[metric_id]
        if definition.get("dependency_only"):
            continue
        name = str(definition["display_name"])
        metricas[name] = [{
            "metric": name, "indicator_id": metric_id,
            "confidence": "medium" if any(item.get("confidence") == "medium" for item in items) else "high",
            "calculated": all(item.get("reported_or_derived") == "derived" for item in items),
            "unit": definition["unit"], "source": "; ".join(dict.fromkeys(str(item.get("source_document") or "") for item in items)),
            "series": {str(item["period"]): item["value"] for item in items},
            "observations": items,
        }]
    return metricas


def build_snapshot(*, ticker: str, company_name: str, observations: list[dict[str, Any]],
                   companies_requested: int, documents_processed: list[str],
                   previous_payload: dict[str, Any] | None = None) -> dict[str, Any]:
    reported = dedupe_observations(observations)
    derived, diagnostics = derive_metrics(reported)
    combined = dedupe_observations([*reported, *derived])
    if not combined and previous_payload and previous_payload.get("observations"):
        preserved = dict(previous_payload)
        preserved["status"] = "preserved_existing_data"
        preserved.setdefault("warnings", []).append("Snapshot anterior preservado: nenhuma observação válida nova encontrada.")
        return preserved
    missing = [
        definition["display_name"] for metric_id, definition in RETAIL_OPERATIONAL_DICTIONARY.items()
        if not definition.get("dependency_only") and not any(item.get("indicator_id") == metric_id for item in combined)
    ]
    return {
        "schema_version": RETAIL_SCHEMA_VERSION, "sector": "varejo",
        "generated_at": datetime.now(timezone.utc).isoformat(), "extractor_version": RETAIL_EXTRACTOR_VERSION,
        "ticker": ticker, "companhia": company_name, "companies_requested": companies_requested,
        "documents_processed": len(documents_processed), "metricas": build_metricas(combined),
        "observations": combined, "status": "found_new_data" if combined else "not_found_no_previous_data",
        "coverage_status": "found" if combined else "not_found",
        "discovery": {"source_policy": "official_spreadsheet_then_official_pdf", "documents_processed": documents_processed},
        "calculation_metadata": {"derived_observations": len(derived), "diagnostics": diagnostics},
        "warnings": ([{
            "metric": name, "status": "not_found",
            "message": "Indicador não divulgado ou sem extração confiável; o valor permanece nulo.",
        } for name in missing]),
    }
